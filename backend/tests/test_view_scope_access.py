"""Tests du périmètre de vue : auto-élargissement (F05) et sous-ressources (F06)."""

import os

import pytest
from app.models import CardComment, User, UserRole, ViewScope
from app.routers.auth import router as auth_router
from app.routers.card_comments import router as card_comments_router
from app.routers.card_items import router as card_items_router
from app.routers.cards import router as cards_router
from app.routers.export import router as export_router
from app.routers.lists import router as lists_router
from app.routers.users import router as users_router
from app.schemas import CardCreate
from app.schemas.card_comment import CardCommentCreate
from app.schemas.card_item import CardItemCreate
from app.services import card as card_service
from app.services import card_comment as card_comment_service
from app.services import card_item as card_item_service

PASSWORD = "Scope123!"
ROUTERS = (
    auth_router,
    users_router,
    cards_router,
    card_comments_router,
    card_items_router,
    lists_router,
    export_router,
)


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def scoped_board(
    integration_session_factory,
    seed_admin_user,
    create_regular_user,
    create_list_record,
):
    """Board avec une carte visible et une carte masquée pour l'utilisateur restreint.

    Renvoie une fonction qui crée l'utilisateur restreint (rôle, périmètre) et les
    données associées.
    """
    seed_admin_user()
    list_id = create_list_record("Backlog", 1)
    other_list_id = create_list_record("Done", 2)
    create_regular_user("owner@example.com", PASSWORD, display_name="Owner")

    def _setup(role: UserRole, view_scope: ViewScope) -> dict:
        email = f"{role.value}@example.com"
        create_regular_user(email, PASSWORD, display_name="Restricted", role=role)
        session = integration_session_factory()
        try:
            owner = session.query(User).filter(User.email == "owner@example.com").one()
            user = session.query(User).filter(User.email == email).one()
            user.view_scope = view_scope
            session.commit()

            hidden = card_service.create_card(
                session,
                CardCreate(
                    title="Carte masquée",
                    description="Secret",
                    list_id=list_id,
                    assignee_id=owner.id,
                ),
                created_by=owner.id,
            )
            visible = card_service.create_card(
                session,
                CardCreate(title="Carte visible", list_id=list_id, assignee_id=user.id),
                created_by=owner.id,
            )
            item = card_item_service.create_item(
                session, CardItemCreate(card_id=hidden.id, text="Tâche secrète")
            )
            comment = card_comment_service.create_comment(
                session,
                CardCommentCreate(card_id=hidden.id, comment="Commentaire secret"),
                owner.id,
            )
            return {
                "email": email,
                "user_id": user.id,
                "list_id": list_id,
                "other_list_id": other_list_id,
                "hidden_id": hidden.id,
                "visible_id": visible.id,
                "item_id": item.id,
                "comment_id": comment.id,
            }
        finally:
            session.close()

    return _setup


# ---------------------------------------------------------------------------
# F05 — modification du périmètre de vue
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_non_admin_cannot_widen_own_view_scope(
    async_client_factory, scoped_board, login_user
):
    data = scoped_board(UserRole.CONTRIBUTOR, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        url = f"/users/{data['user_id']}/view-scope"

        for wider in ("all", "unassigned_plus_mine"):
            response = await client.put(
                url, json={"view_scope": wider}, headers=_auth(token)
            )
            assert response.status_code == 403

        # Conserver le même périmètre reste permis
        response = await client.put(
            url, json={"view_scope": "mine_only"}, headers=_auth(token)
        )
        assert response.status_code == 200

        # Toujours aucune carte d'autrui visible
        response = await client.get(f"/cards/{data['hidden_id']}", headers=_auth(token))
        assert response.status_code == 403


@pytest.mark.asyncio
async def test_non_admin_can_reduce_own_view_scope(
    async_client_factory, scoped_board, login_user
):
    data = scoped_board(UserRole.EDITOR, ViewScope.ALL)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        url = f"/users/{data['user_id']}/view-scope"

        response = await client.put(
            url, json={"view_scope": "unassigned_plus_mine"}, headers=_auth(token)
        )
        assert response.status_code == 200
        assert response.json()["view_scope"] == "unassigned_plus_mine"

        response = await client.put(
            url, json={"view_scope": "mine_only"}, headers=_auth(token)
        )
        assert response.status_code == 200

        # Revenir en arrière est un élargissement : refusé
        response = await client.put(
            url, json={"view_scope": "all"}, headers=_auth(token)
        )
        assert response.status_code == 403


@pytest.mark.asyncio
async def test_admin_can_widen_view_scope(
    async_client_factory, scoped_board, login_user
):
    data = scoped_board(UserRole.CONTRIBUTOR, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, "admin@yaka.local", "Admin-Test1")
        response = await client.put(
            f"/users/{data['user_id']}/view-scope",
            json={"view_scope": "all"},
            headers=_auth(token),
        )
        assert response.status_code == 200
        assert response.json()["view_scope"] == "all"


# ---------------------------------------------------------------------------
# F06 — périmètre appliqué aux sous-ressources, agrégats et mutations
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_visitor_mine_only_cannot_read_hidden_card_resources(
    async_client_factory, scoped_board, login_user
):
    data = scoped_board(UserRole.VISITOR, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        headers = _auth(token)
        hidden = data["hidden_id"]

        for url in (
            f"/cards/{hidden}",
            f"/cards/{hidden}/history",
            f"/card-comments/card/{hidden}",
            f"/card-items/card/{hidden}",
        ):
            response = await client.get(url, headers=headers)
            assert response.status_code == 403, url
            assert "secret" not in response.text.lower()

        # La carte visible reste accessible avec ses sous-ressources
        visible = data["visible_id"]
        for url in (
            f"/cards/{visible}",
            f"/cards/{visible}/history",
            f"/card-comments/card/{visible}",
            f"/card-items/card/{visible}",
        ):
            response = await client.get(url, headers=headers)
            assert response.status_code == 200, url

        # Une carte inexistante reste une 404
        response = await client.get("/card-items/card/999999", headers=headers)
        assert response.status_code == 404

        # Le compteur de la liste ne compte que les cartes visibles
        response = await client.get(
            f"/lists/{data['list_id']}/cards-count", headers=headers
        )
        assert response.status_code == 200
        assert response.json()["cards_count"] == 1


@pytest.mark.asyncio
async def test_export_is_filtered_by_view_scope(
    async_client_factory, scoped_board, login_user
):
    data = scoped_board(UserRole.VISITOR, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        response = await client.get("/export/?format=csv", headers=_auth(token))
        assert response.status_code == 200
        content = response.content.decode("utf-8-sig")
        assert "Carte visible" in content
        assert "Carte masquée" not in content

        response = await client.get("/export/?format=xlsx", headers=_auth(token))
        assert response.status_code == 200

    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(response.content)).active
    titles = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert titles == ["Carte visible"]


@pytest.mark.asyncio
@pytest.mark.parametrize("role", [UserRole.CONTRIBUTOR, UserRole.SUPERVISOR])
async def test_mine_only_cannot_mutate_hidden_card(
    async_client_factory, scoped_board, login_user, role
):
    # SUPERVISOR peut modifier toutes les cartes : seul le périmètre le bloque
    data = scoped_board(role, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        headers = _auth(token)
        hidden = data["hidden_id"]

        requests = [
            ("put", f"/cards/{hidden}", {"title": "Piratée"}),
            ("patch", f"/cards/{hidden}/list", {"list_id": data["other_list_id"]}),
            (
                "patch",
                f"/cards/{hidden}/move",
                {
                    "source_list_id": data["list_id"],
                    "target_list_id": data["other_list_id"],
                },
            ),
            ("patch", f"/cards/{hidden}/archive", None),
            ("patch", f"/cards/{hidden}/unarchive", None),
            ("patch", f"/cards/{hidden}/statut?statut=termine", None),
            (
                "post",
                f"/cards/{hidden}/history",
                {
                    "card_id": hidden,
                    "user_id": data["user_id"],
                    "action": "x",
                    "description": "x",
                },
            ),
            ("delete", f"/cards/{hidden}", None),
            (
                "post",
                "/cards/bulk-move",
                {"card_ids": [hidden], "target_list_id": data["other_list_id"]},
            ),
            ("post", "/card-comments/", {"card_id": hidden, "comment": "Coucou"}),
            ("post", "/card-items/", {"card_id": hidden, "text": "Nouvelle"}),
            ("put", f"/card-items/{data['item_id']}", {"is_done": True}),
            ("delete", f"/card-items/{data['item_id']}", None),
            ("put", f"/card-comments/{data['comment_id']}", {"comment": "Modifié"}),
            ("delete", f"/card-comments/{data['comment_id']}", None),
        ]
        for method, url, body in requests:
            kwargs = {"headers": headers}
            if body is not None:
                kwargs["json"] = body
            response = await client.request(method.upper(), url, **kwargs)
            assert response.status_code == 403, (method, url, response.text)
            assert "Secret" not in response.text


@pytest.mark.asyncio
async def test_supervisor_mine_only_does_not_get_hidden_card_via_archive(
    async_client_factory, scoped_board, login_user, integration_session_factory
):
    data = scoped_board(UserRole.SUPERVISOR, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        response = await client.patch(
            f"/cards/{data['hidden_id']}/archive", headers=_auth(token)
        )
        assert response.status_code == 403
        assert "Carte masquée" not in response.text
        assert "Secret" not in response.text

    session = integration_session_factory()
    try:
        card = card_service.get_card(session, data["hidden_id"])
        assert card is not None and card.is_archived is False
    finally:
        session.close()


@pytest.mark.asyncio
async def test_admin_restricted_scope_counts_and_exports_all_cards(
    async_client_factory, scoped_board, login_user, integration_session_factory
):
    # Le décompte sert à confirmer la suppression d'une liste (admin) : il ne
    # doit pas dépendre du périmètre d'affichage de l'admin.
    data = scoped_board(UserRole.CONTRIBUTOR, ViewScope.ALL)
    session = integration_session_factory()
    try:
        admin = session.query(User).filter(User.role == UserRole.ADMIN).first()
        admin.view_scope = ViewScope.MINE_ONLY
        admin_email = admin.email
        session.commit()
    finally:
        session.close()

    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(
            client, admin_email, os.environ["DEFAULT_ADMIN_PASSWORD"]
        )
        headers = _auth(token)
        response = await client.get(
            f"/lists/{data['list_id']}/cards-count", headers=headers
        )
        assert response.status_code == 200
        assert response.json()["cards_count"] == 2

        response = await client.get("/export/?format=csv", headers=headers)
        content = response.content.decode("utf-8-sig")
        assert "Carte visible" in content and "Carte masquée" in content


@pytest.mark.asyncio
async def test_history_entry_body_card_id_must_match_url(
    async_client_factory, scoped_board, login_user
):
    data = scoped_board(UserRole.SUPERVISOR, ViewScope.MINE_ONLY)
    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        response = await client.post(
            f"/cards/{data['visible_id']}/history",
            json={
                "card_id": data["hidden_id"],
                "user_id": data["user_id"],
                "action": "x",
                "description": "x",
            },
            headers=_auth(token),
        )
        assert response.status_code == 400


def test_get_card_does_not_purge_soft_deleted_comments(
    integration_session_factory, scoped_board
):
    """Régression : get_card suivi d'un commit ne doit pas supprimer les
    commentaires soft-deleted (cascade delete-orphan sur Card.comments)."""
    data = scoped_board(UserRole.CONTRIBUTOR, ViewScope.ALL)
    session = integration_session_factory()
    try:
        owner = session.query(User).filter(User.email == "owner@example.com").one()
        deleted = card_comment_service.create_comment(
            session,
            CardCommentCreate(card_id=data["hidden_id"], comment="À supprimer"),
            owner.id,
        )
        assert card_comment_service.delete_comment(session, deleted.id, owner.id)

        card = card_service.get_card(session, data["hidden_id"])
        assert [c.id for c in card.comments] == [data["comment_id"]]
        session.commit()
        session.expire_all()

        ids = {
            c.id
            for c in session.query(CardComment)
            .filter(CardComment.card_id == data["hidden_id"])
            .all()
        }
        assert ids == {data["comment_id"], deleted.id}
    finally:
        session.close()


def _soft_delete_comment(session, card_id: int) -> int:
    """Créer puis supprimer logiquement un commentaire du propriétaire."""
    owner = session.query(User).filter(User.email == "owner@example.com").one()
    deleted = card_comment_service.create_comment(
        session, CardCommentCreate(card_id=card_id, comment="À supprimer"), owner.id
    )
    assert card_comment_service.delete_comment(session, deleted.id, owner.id)
    return deleted.id


@pytest.mark.asyncio
async def test_card_mutation_response_hides_soft_deleted_comments(
    async_client_factory, integration_session_factory, scoped_board, login_user
):
    """Régression : après commit + refresh, la réponse d'une mutation ne doit pas
    réexposer les commentaires soft-deleted."""
    data = scoped_board(UserRole.SUPERVISOR, ViewScope.ALL)
    session = integration_session_factory()
    try:
        deleted_id = _soft_delete_comment(session, data["hidden_id"])
    finally:
        session.close()

    async with async_client_factory(*ROUTERS) as client:
        token = await login_user(client, data["email"], PASSWORD)
        response = await client.put(
            f"/cards/{data['hidden_id']}",
            json={"title": "Titre modifié"},
            headers=_auth(token),
        )
        assert response.status_code == 200
        ids = [c["id"] for c in response.json()["comments"]]
        assert ids == [data["comment_id"]]
        assert deleted_id not in ids


def test_delete_card_removes_soft_deleted_comments(
    integration_session_factory, scoped_board
):
    """Régression : la suppression d'une carte doit aussi supprimer ses
    commentaires soft-deleted (masqués par get_card, pas de FK SQLite)."""
    data = scoped_board(UserRole.CONTRIBUTOR, ViewScope.ALL)
    session = integration_session_factory()
    try:
        _soft_delete_comment(session, data["hidden_id"])
        assert card_service.delete_card(session, data["hidden_id"])
        session.expire_all()

        remaining = (
            session.query(CardComment)
            .filter(CardComment.card_id == data["hidden_id"])
            .count()
        )
        assert remaining == 0
    finally:
        session.close()
